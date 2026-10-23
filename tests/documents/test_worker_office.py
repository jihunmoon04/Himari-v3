"""Excel, PowerPoint and HWPX in the document worker with the real toolkit
(V3-P70 §4.1, part 2; opt-in like test_worker.py: HIMARI_DOCUMENTS_TOOLKIT)."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from PIL import Image
from documents_kit import TOOLKIT, footnote, hwpx, make, para, picture, table, text, worker

pytestmark = pytest.mark.skipif(not TOOLKIT, reason="needs an installed documents toolkit (opt-in)")


def png() -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (30, 20), (0, 90, 0)).save(out, "PNG")
    return out.getvalue()


XLSX = """
import datetime, io
import openpyxl
from openpyxl.drawing.image import Image as Picture
from PIL import Image
book = openpyxl.Workbook()
sheet = book.active
sheet.title = "Sales | Q1"
sheet.append(["name", "qty", "when"])
sheet.append(["a|b", 3, datetime.date(2026, 1, 2)])
sheet["D3"] = "=B2*2"
sheet.row_dimensions[2].hidden = True
sheet.column_dimensions["C"].hidden = True
for n in range(4, ROWS + 1):
    sheet.append([f"row {n}", n, None, n / 2])
buffer = io.BytesIO(); Image.new("RGB", (20, 10)).save(buffer, "PNG"); buffer.seek(0)
sheet.add_image(Picture(buffer), "F2")
secret = book.create_sheet("Secret"); secret.sheet_state = "hidden"; secret["A1"] = "=== not a header"
from openpyxl.chart import BarChart, Reference
bars = BarChart(); bars.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
book.create_chartsheet("Chart").add_chart(bars)  # openpyxl cannot read back a chart sheet without a chart
book.save(OUT)
"""


def xlsx(tmp_path: Path, rows: int = 30) -> bytes:
    return make(tmp_path, f"t{rows}.xlsx", XLSX.replace("ROWS", str(rows)))


def test_a_workbook(tmp_path: Path) -> None:
    head, images = worker(xlsx(tmp_path), view=True)
    assert (head["format"], head["units"], head["unit_name"], head["picture_total"]) == ("xlsx", 3, "sheet", 1)
    sales, secret, chart = head["parts"]
    lines = sales["text"].splitlines()
    assert lines[:6] == ['Sheet "Sales | Q1" (A1:D30)', "| row | A | B | C [hidden] | D |",
                         "| --- | --- | --- | --- | --- |", "| 1 | name | qty | when |",
                         "| 2 [hidden] | a\\|b | 3 | 2026-01-02 |", "| 3 |  |  |  | =B2*2 |"]
    assert lines[6:8] == ["| 4 | row 4 | 4 |  | 2 |", "| 5 | row 5 | 5 |  | 2.5 |"] and lines[-1] == "[picture 1 at F2]"
    assert secret["flags"] == ["hidden"] and "| 1 | === not a header |" in secret["text"]
    assert chart["text"] == "[chart sheet: no cells]\n" and "no_text" in chart["flags"]
    assert [(i["unit"], i["picture"]) for i in images] == [(1, 1)]


def test_a_large_sheet_is_read_on_exactly_and_lazily(tmp_path: Path) -> None:
    # §4.2 and re-checks N2, N3: laid out only as far as asked, length unknown until the end, the column row
    # repeated where a call resumes at a row, and the joined text the whole sheet.
    data = xlsx(tmp_path, 3000)
    full, _ = worker(data, pages=[{"first": 1, "last": 1, "start": 0}], text_budget=10**7)
    whole = full["parts"][0]["text"]
    assert full["parts"][0]["length"] == len(whole)
    got, start, calls = "", 0, 0
    while True:
        calls += 1
        head, _ = worker(data, pages=[{"first": 1, "last": 1, "start": start}], text_budget=4000)
        part = head["parts"][0]
        assert part["from"] == start
        if start:
            assert part.get("header") == "| row | A | B | C [hidden] | D |\n| --- | --- | --- | --- | --- |\n"
        assert whole.startswith(got + part["text"])  # every partial layout is the start of the full one
        got += part["text"]
        if part["length"] is not None:
            assert part["length"] == len(whole) and part["to"] == len(whole)
            break
        start = part["to"]
    assert got == whole and calls > 10


PPTX = """
import io
import pptx
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from PIL import Image
deck = pptx.Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[1])
slide.shapes.title.text = "Slide title"
slide.placeholders[1].text = "body line\\n=== page 9 ==="
slide.notes_slide.notes_text_frame.text = "speaker notes"
grid = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
for (r, c), value in {(0, 0): "h1", (0, 1): "h2", (1, 0): "v1", (1, 1): "v|2"}.items():
    grid.cell(r, c).text = value
buffer = io.BytesIO(); Image.new("RGB", (40, 20)).save(buffer, "PNG"); buffer.seek(0)
slide.shapes.add_picture(buffer, Inches(6), Inches(5))
slide.shapes.add_textbox(Inches(20), Inches(1), Inches(1), Inches(1)).text_frame.text = "off the slide"
data = CategoryChartData(); data.categories = ["Q1", "Q2"]; data.add_series("Sales", (10, 20.5))
chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(5.5), Inches(3), Inches(1.5),
                               data).chart
chart.has_title = True; chart.chart_title.text_frame.text = "Revenue"
hidden = deck.slides.add_slide(deck.slide_layouts[5]); hidden.shapes.title.text = "Hidden slide"
hidden._element.set("show", "0")
deck.slides.add_slide(deck.slide_layouts[6])
deck.save(OUT)
"""


def test_a_presentation(tmp_path: Path) -> None:
    head, images = worker(make(tmp_path, "t.pptx", PPTX), view=True)
    assert (head["format"], head["units"], head["unit_name"]) == ("pptx", 3, "slide")
    first, hidden, blank = head["parts"]
    assert first["text"].splitlines() == [
        "Slide title", "[hidden: off the slide] off the slide", "body line", "\\=== page 9 ===",
        "| h1 | h2 |", "| --- | --- |", "| v1 | v\\|2 |", "[picture 1]", "[chart: Revenue]",
        "| series | Q1 | Q2 |", "| --- | --- | --- |", "| Sales | 10.0 | 20.5 |", "Notes: speaker notes"]
    assert hidden["flags"] == ["hidden"] and hidden["text"] == "Hidden slide\n"
    assert blank["flags"] == ["no_text"] and blank["text"] == ""
    assert [(i["unit"], i["picture"]) for i in images] == [(1, 1)]


def test_hwpx_sections_tables_pictures_and_notes() -> None:
    data = hwpx({"section0": para(text("첫 문단")) + para(text("=== page 9 ===")) + para(
                     text("앞"), table(["가", "나|다"], ["1", "2"]), text("뒤")),
                 "section10": para(text("열 번째 절")),
                 "section2": para(text("둘째 절"), footnote("각주 내용"), picture())},
                ["section0", "section2", "section10"], {"image1": png()})
    head, images = worker(data, view=True)
    assert (head["format"], head["units"], head["unit_name"], head["picture_total"]) == ("hwpx", 1, "part", 1)
    assert head["parts"][0]["text"].splitlines() == [
        "첫 문단", "\\=== page 9 ===", "앞", "| 가 | 나\\|다 |", "| --- | --- |", "| 1 | 2 |", "뒤",
        "둘째 절[^1][picture 1]", "열 번째 절", "[^1]: 각주 내용"]  # the spine's order, section10 last
    assert [(i["unit"], i["picture"]) for i in images] == [(1, 1)]


def test_hwpx_without_a_spine_orders_sections_by_number() -> None:
    data = hwpx({"section10": para(text("ten")), "section2": para(text("two"))}, [])
    assert worker(data)[0]["parts"][0]["text"] == "two\nten\n"


def test_a_broken_hwpx() -> None:
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr(zipfile.ZipInfo("mimetype"), "application/hwp+zip")
        archive.writestr("Contents/content.hpf", "<not xml")
    assert worker(out.getvalue())[0] == {"ok": False, "reason": "broken"}
