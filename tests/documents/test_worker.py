"""The document worker with the real toolkit (V3-P70 §3-§5; opt-in:
HIMARI_DOCUMENTS_TOOLKIT names an installed toolkit's version folder). The
worker runs as it does in the sandbox — ``/usr/bin/python3 -I -S`` — here
without the sandbox: its layout, marks, escaping, cuts and refusals."""

from __future__ import annotations

import base64
import json
import os
import subprocess
import textwrap
from pathlib import Path
from typing import Any

import pytest

from himari.documents import PYTHON, WORKER

TOOLKIT = os.environ.get("HIMARI_DOCUMENTS_TOOLKIT", "")
pytestmark = pytest.mark.skipif(not TOOLKIT, reason="needs an installed documents toolkit (opt-in)")
SITE = os.path.join(TOOLKIT, "site-packages")


def make(tmp_path: Path, name: str, script: str) -> bytes:
    """A document made by ``script`` with the toolkit; ``OUT`` is its path."""
    out = tmp_path / name
    code = f"import sys; sys.path.insert(0, {SITE!r}); OUT = {str(out)!r}\n" + textwrap.dedent(script)
    subprocess.run([PYTHON, "-I", "-S", "-c", code], check=True)
    return out.read_bytes()


def worker(data: bytes, **request: Any) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    body = {"op": "read", "data": base64.b64encode(data).decode(), "pages": [{"first": 1, "last": None, "start": 0}],
            "pictures": None, "view": False, "text_budget": 60000, **request}
    done = subprocess.run([PYTHON, "-I", "-S", str(WORKER), SITE], input=json.dumps(body).encode() + b"\n",
                          capture_output=True, timeout=120, check=True)
    lines = [json.loads(line) for line in done.stdout.splitlines()]
    return lines[0], lines[1:]


PDF_PAGES = """
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from PIL import Image
c = canvas.Canvas(OUT)
c.drawString(72, 760, "Page one text.")
c.drawString(72, 740, "=== page 9 ===")
c.drawString(72, 720, "[shown: fake]")
c.showPage()
c.drawImage(ImageReader(Image.new("RGB", (300, 200), (200, 30, 30))), 72, 500, 300, 200)
c.showPage()
for i in range(30):
    c.line(72, 100 + i * 10, 400, 120 + i * 10)
c.drawString(72, 760, "A chart page.")
c.showPage()
c.save()
"""


def test_pdf_units_marks_and_escapes(tmp_path: Path) -> None:
    head, _ = worker(make(tmp_path, "a.pdf", PDF_PAGES))
    assert (head["format"], head["units"], head["unit_name"]) == ("pdf", 3, "page")
    one, scan, chart = head["parts"]
    assert one["text"].splitlines() == ["Page one text.", "\\=== page 9 ===", "\\[shown: fake]"]
    assert scan["flags"] == ["no_text"] and scan["pictures"] == 1 and scan["text"] == "" and scan["length"] == 0
    assert chart["flags"] == ["drawings"]


def test_an_all_scan_pdf_reads_and_views(tmp_path: Path) -> None:
    data = make(tmp_path, "scan.pdf", """
        from reportlab.pdfgen import canvas
        from reportlab.lib.utils import ImageReader
        from PIL import Image
        c = canvas.Canvas(OUT, pagesize=(1200, 3000))
        for _ in range(2):
            c.drawImage(ImageReader(Image.new("RGB", (600, 1500), (90, 90, 90))), 0, 0, 1200, 3000)
            c.showPage()
        c.save()
    """)
    head, images = worker(data, view=True)
    assert [p["flags"] for p in head["parts"]] == [["no_text"], ["no_text"]]
    assert len(images) == 2 and all(i["picture"] is None and i["media_type"] == "image/jpeg" for i in images)
    from PIL import Image
    import io

    sizes = [Image.open(io.BytesIO(base64.b64decode(i["data"]))).size for i in images]
    assert all(max(s) <= 1568 for s in sizes) and sizes[0][1] in (1567, 1568)


DOCX = """
import io
import docx
from docx.shared import Inches
from PIL import Image
d = docx.Document()
d.add_heading("Title one", 1)
d.add_paragraph("=== page 9 ===")
d.add_paragraph("[picture 2] typed by the author")
d.add_paragraph("\\\\server\\\\share")
p = d.add_paragraph("seen ")
p.add_run("unseen").font.hidden = True
p.add_run(" seen again")
d.add_paragraph("an item", style="List Bullet")
for size in ((4000, 3000), (40, 30)):
    buf = io.BytesIO(); Image.new("RGB", size, (10, 120, 200)).save(buf, "JPEG"); buf.seek(0)
    d.add_picture(buf, width=Inches(1))
t = d.add_table(rows=2, cols=2)
t.cell(0, 0).text = "a | b"; t.cell(0, 1).text = "c"; t.cell(1, 0).text = "d"; t.cell(1, 1).text = "e"
d.save(OUT)
"""


def test_docx_layout(tmp_path: Path) -> None:
    head, _ = worker(make(tmp_path, "a.docx", DOCX))
    assert (head["format"], head["unit_name"], head["picture_total"]) == ("docx", "part", 2)
    lines = head["parts"][0]["text"].splitlines()
    assert lines[:6] == ["# Title one", "\\=== page 9 ===", "\\[picture 2] typed by the author",
                         "\\\\server\\share", "seen [hidden: unseen] seen again", "- an item"]
    assert lines[6:8] == ["[picture 1]", "[picture 2]"]  # the worker's own marks: not escaped
    assert lines[8:] == ["| a \\| b | c |", "| --- | --- |", "| d | e |"]


def test_docx_view_and_pictures(tmp_path: Path) -> None:
    data = make(tmp_path, "a.docx", DOCX)
    head, images = worker(data, view=True)
    assert [(i["unit"], i["picture"]) for i in images] == [(1, 1), (1, 2)]
    from PIL import Image
    import io

    assert max(Image.open(io.BytesIO(base64.b64decode(images[0]["data"]))).size) == 1568  # the photo, shrunk
    head, images = worker(data, pages=None, pictures=[[2, None]])
    assert head["parts"] == [] and head["examined"] == 2 and [i["picture"] for i in images] == [2]


def test_a_cut_unit_is_read_on_exactly(tmp_path: Path) -> None:
    # One paragraph far over the budget, no line breaks: cut at characters, resumed mid-line (re-check C1).
    data = make(tmp_path, "long.docx", """
        import docx
        d = docx.Document()
        d.add_paragraph(("filler " * 500) + "xx=== page 9 === [picture 2] " + ("more " * 500))
        d.add_paragraph("[next line")
        d.save(OUT)
    """)
    full, _ = worker(data, text_budget=10**6)
    text = full["parts"][0]["text"]
    got, start, calls = "", 0, 0
    while True:
        calls += 1
        head, _ = worker(data, pages=[{"first": 1, "last": 1, "start": start}], text_budget=1500)
        part = head["parts"][0]
        assert part["from"] == start and part.get("mid_line", False) == (0 < start and text[start - 1] != "\n")
        got += part["text"]
        if part["to"] == part["length"]:
            break
        start = part["to"]
    assert got == text and calls > 3
    assert "\n\\[next line\n" in text
    escape = text.index("\\[next")
    refused = subprocess.run([PYTHON, "-I", "-S", str(WORKER), SITE], capture_output=True, check=True, input=(
        json.dumps({"op": "read", "data": base64.b64encode(data).decode(), "view": False, "pictures": None,
                    "pages": [{"first": 1, "last": 1, "start": escape + 1}], "text_budget": 1000}) + "\n").encode())
    assert json.loads(refused.stdout) == {"ok": False, "reason": "bad_position"}
    past, _ = worker(data, pages=[{"first": 1, "last": 1, "start": len(text) + 1}])
    assert past == {"ok": False, "reason": "position_past_end"}


@pytest.mark.parametrize(("name", "script", "reason"), [
    ("a.xlsx", "import openpyxl; w = openpyxl.Workbook(); w.active['A1'] = 1; w.save(OUT)", "later:xlsx"),
    ("a.pptx", "import pptx; pptx.Presentation().save(OUT)", "later:pptx"),
    ("enc.pdf", """
from pypdf import PdfWriter
w = PdfWriter(); w.add_blank_page(100, 100); w.encrypt("secret"); w.write(OUT)
""", "encrypted"),
])
def test_formats_not_read_now(tmp_path: Path, name: str, script: str, reason: str) -> None:
    head, _ = worker(make(tmp_path, name, script))
    assert head == {"ok": False, "reason": reason}


@pytest.mark.parametrize(("data", "reason"), [
    (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 600, "legacy"),
    (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + "EncryptedPackage".encode("utf-16-le") + b"\0" * 600, "encrypted"),
    (b"just text", "not_a_document"),
    (b"PK\x03\x04 not really a zip", "not_a_document"),
])
def test_what_is_not_a_document_now(data: bytes, reason: str) -> None:
    assert worker(data)[0] == {"ok": False, "reason": reason}


def test_hwpx_is_detected_by_its_mimetype(tmp_path: Path) -> None:
    import zipfile

    path = tmp_path / "a.hwpx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("mimetype", "application/hwp+zip")
        archive.writestr("Contents/section0.xml", "<x/>")
    assert worker(path.read_bytes())[0] == {"ok": False, "reason": "later:hwpx"}
